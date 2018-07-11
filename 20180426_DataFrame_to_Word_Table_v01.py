import pandas as pd
import docx
import numpy as np
from dump_df_to_word import *

doc_name = 'Test.docx'
default_path = './'

df = pd.DataFrame(np.arange(100).reshape((25,4)),columns=['a','b','c','d'])

dump_df_to_word(df,doc_name,default_path)

'''
doc = docx.Document()

t = doc.add_table(df.shape[0]+1,df.shape[1])

for j in range(df.shape[-1]):
    t.cell(0,j).text = df.columns[j]

for i in range(df.shape[0]):
    for j in range(df.shape[-1]):
        t.cell(i+1,j).text = str(df.values[i,j])

doc.save(default_path+doc_name)
'''
